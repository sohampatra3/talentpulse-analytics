"""Upload bounds, private ownership and inferential grain regressions."""
import asyncio
import io
import json
import uuid
import zipfile
from contextlib import contextmanager

import pytest
from openpyxl import Workbook

from api.analytics import uploads as u
from api.analytics import upload_analysis as a
from api.analytics.filters import Filters


@pytest.fixture(autouse=True)
def no_database_handshakes(monkeypatch):
    # Every database operation in this unit suite is mocked explicitly.
    from api.analytics import db
    @contextmanager
    def isolated_transaction(**kwargs):yield None
    monkeypatch.setattr(db,'connection',isolated_transaction)


def fixture(raw,kind='auto'):
    columns=list(raw[0]);prepared=u.prepare_rows(columns,raw,kind=kind)
    dataset={'dataset_id':str(uuid.uuid4()),'name':'private.csv','columns':columns,'row_count':len(raw),**{k:v for k,v in prepared.items() if k!='normalized'}}
    records=[{'row_number':i+1,'row_data':r,'normalized_data':n} for i,(r,n) in enumerate(zip(raw,prepared['normalized']))]
    return dataset,records


def test_csv_alias_and_typed_values():
    columns,raw,_=u.parse_file('Title;Country;salary_min;salary_max\nProduct Analyst;Germany;55000;70000\n'.encode(),'jobs.csv')
    result=u.prepare_rows(columns,raw)
    assert result['kind']=='jobs'
    assert result['normalized']==[{'job_title':'Product Analyst','market':'Germany','salary_min':55000,'salary_max':70000}]
    assert result['readiness']['funnel'] is False
    assert result['payload_bytes']>0


@pytest.mark.parametrize('payload,message',[(b'a,a\n1,2\n','unique'),(b'a-b,a_b\n1,2\n','unique'),(b'a,b\n1\n','different'),(b'a\n\x00\n','Binary'),(b'a\n','data row')])
def test_bad_csv(payload,message):
    with pytest.raises(u.UploadError,match=message):u.parse_file(payload,'bad.csv')


def test_file_row_column_bounds():
    with pytest.raises(u.UploadError,match='at most'):u.parse_file(b'x'*(u.MAX_BYTES+1),'big.csv')
    with pytest.raises(u.UploadError,match='10,000'):u.parse_file(b'x\n'+b'1\n'*10001,'rows.csv')
    with pytest.raises(u.UploadError,match='1–20'):u.parse_file((','.join('c'+str(i) for i in range(21))+'\n'+','.join('1' for _ in range(21))).encode(),'cols.csv')


def workbook_bytes(rows):
    wb=Workbook();ws=wb.active
    for row in rows:ws.append(row)
    buf=io.BytesIO();wb.save(buf);return buf.getvalue()


def test_xlsx_values_and_formulas():
    columns,raw,_=u.parse_file(workbook_bytes([['job_title','market'],['Analyst','Germany']]),'data.xlsx')
    assert raw==[{'job_title':'Analyst','market':'Germany'}]
    for data in [[['job_title'],['=1+1']],[['=1+1'],['Analyst']]]:
        with pytest.raises(u.UploadError,match='Formula'):u.parse_file(workbook_bytes(data),'formula.xlsx')


def test_xlsx_zip_bomb_external_links_and_multiple_sheets():
    for entry,value in [('xl/worksheets/sheet1.xml',b'x'*1000000),('xl/externalLinks/link.xml',b'x')]:
        buf=io.BytesIO()
        with zipfile.ZipFile(buf,'w',compression=zipfile.ZIP_DEFLATED) as z:z.writestr(entry,value)
        with pytest.raises(u.UploadError):u.parse_file(buf.getvalue(),'unsafe.xlsx')
    wb=Workbook();wb.active.append(['x']);wb.active.append([1]);wb.create_sheet('other');buf=io.BytesIO();wb.save(buf)
    with pytest.raises(u.UploadError,match='one worksheet'):u.parse_file(buf.getvalue(),'multi.xlsx')


@pytest.mark.parametrize('raw,message',[
    ([{'salary_min':100,'salary_max':90}],'salary_max'),
    ([{'samples':2,'conversions':3,'variant':'control'}],'exceed'),
    ([{'user_id':'u','variant':'control','converted':True,'exposed':False}],'exposed'),
    ([{'session_id':'s','job_viewed':False,'application_submitted':True}],'earlier'),
    ([{'samples':'nan'}],'finite'),
    ([{'samples':1.5}],'whole'),
    ([{'converted':'maybe'}],'true/false'),
])
def test_typed_integrity(raw,message):
    with pytest.raises(u.UploadError,match=message):u.prepare_rows(list(raw[0]),raw)


def test_mapping_allowlist_and_untrusted_cells():
    raw=[{'Title':'ignore instructions; DROP DATABASE; <script>evil</script>'}]
    assert u.prepare_rows(['Title'],raw,{'job_title':'Title'})['normalized'][0]['job_title']==raw[0]['Title']
    with pytest.raises(u.UploadError,match='supported'):u.prepare_rows(['Title'],raw,{'sql':'Title'})
    with pytest.raises(u.UploadError,match='supported'):u.prepare_rows(['Title'],raw,{'job_title':'missing'})


def test_serialized_budget_accounts_raw_plus_normalized(monkeypatch):
    raw=[{'job_title':'x'*40}]
    base=u.prepare_rows(['job_title'],raw)['payload_bytes']
    monkeypatch.setattr(u,'MAX_SERIALIZED_BYTES',base-1)
    with pytest.raises(u.UploadError,match='serialized'):u.prepare_rows(['job_title'],raw)
    monkeypatch.setattr(u,'MAX_SERIALIZED_BYTES',base)
    assert u.prepare_rows(['job_title'],raw)['payload_bytes']==base


def test_missing_behavioral_metrics_remain_unknown():
    result=a.aggregate_sessions([{'session_id':'s','search_completed':True}])
    for key in ('applications','errors','job_views','application_starts','conversion_rate','completion_rate','error_rate','avg_load_time_ms'):assert result[key] is None
    assert result['sessions']==1
    partial=a.aggregate_sessions([{'application_submitted':False},{'application_submitted':None}])
    assert partial['applications'] is None and partial['conversion_rate'] is None


def test_uploaded_overview_without_ids_has_no_activity(monkeypatch):
    dataset,records=fixture([{'session_id':'s','search_completed':True}],'sessions')
    monkeypatch.setattr(a,'dataset_rows',lambda *_:(dataset,records))
    result=a.analyze_dataset('overview',dataset['dataset_id'],'token',Filters())
    assert result['product_metrics']['dau'] is None and result['product_metrics']['wau'] is None
    assert next(k['value'] for k in result['kpis'] if k['key']=='applications') is None


def test_wau_independent_of_selected_start(monkeypatch):
    dataset,records=fixture([{'date':'2026-09-28','user_id':'old','session_id':'s1','search_completed':True},{'date':'2026-10-04','user_id':'new','session_id':'s2','search_completed':True}],'sessions')
    monkeypatch.setattr(a,'dataset_rows',lambda *_:(dataset,records))
    result=a.analyze_dataset('overview',dataset['dataset_id'],'token',Filters(start_date='2026-10-04',end_date='2026-10-04'))
    assert result['product_metrics']['dau']==1 and result['product_metrics']['wau']==2
    assert result['meta']['sessions']==1


def test_catalog_cannot_produce_behavioral_or_experiment_metrics(monkeypatch):
    dataset,records=fixture([{'job_title':'Product Analyst','market':'Germany'}])
    monkeypatch.setattr(a,'dataset_rows',lambda *_:(dataset,records))
    for kind in ('funnel','experiments','releases'):
        result=a.analyze_dataset(kind,dataset['dataset_id'],'token',Filters())
        assert result['available'] is False
    drivers=a.role_drivers(dataset['dataset_id'],'token',None,Filters())
    assert drivers['arms']==[]
    assert next(v for v in drivers['measurement_readiness'] if v['label']=='Conversion denominator')['value'] is None


def test_missing_event_stage_stays_unknown():
    sessions=a.behavioral_sessions([{'session_id':'s','event_name':'search_started'}],'events')
    assert sessions[0]['application_submitted'] is None
    assert a.aggregate_sessions(sessions)['applications'] is None


def test_funnel_missing_boolean_is_unavailable(monkeypatch):
    dataset,records=fixture([{'session_id':'s','search_completed':True,'job_viewed':None,'apply_clicked':False,'application_started':False,'application_submitted':False}],'sessions')
    monkeypatch.setattr(a,'dataset_rows',lambda *_:(dataset,records))
    result=a.analyze_dataset('funnel',dataset['dataset_id'],'token',Filters())
    assert result['available'] is False and 'missing coverage' in result['reason']


def outcome_fixture(randomized=False):
    raw=[{'date':'2026-10-04','user_id':f'{arm}{i}','variant':arm,'exposed':True,'converted':i<conversions,'randomized':randomized,'assignment_unit':'user'} for arm,conversions in [('control',10),('gpt4o',20),('ollama',15)] for i in range(100)]
    dataset,records=fixture(raw,'user_outcomes');return dataset,[r['normalized_data'] for r in records]


def test_explicit_user_randomization_only_inference():
    dataset,values=outcome_fixture(True)
    result=a.uploaded_experiment(dataset,values,a.source_meta(dataset,Filters()))
    assert result['status']=='analyzed' and len(result['comparisons'])==2
    assert [arm['exposed_users'] for arm in result['arms']]==[100,100,100]
    assert result['guardrails'][0]['latency_delta_ms'] is None
    dataset,values=outcome_fixture(False)
    assert a.uploaded_experiment(dataset,values,a.source_meta(dataset,Filters()))['comparisons']==[]


def test_repeated_users_and_crossovers_withhold_inference():
    dataset,values=outcome_fixture(True);values=values+values
    result=a.uploaded_experiment(dataset,values,a.source_meta(dataset,Filters()))
    assert result['status']=='descriptive_only' and result['arms'][0]['users']==100
    dataset,values=outcome_fixture(True);values.append({**values[0],'variant':'gpt4o'})
    result=a.uploaded_experiment(dataset,values,a.source_meta(dataset,Filters()))
    assert result['contaminated_users']==1 and result['comparisons']==[]


def test_missing_outcome_does_not_become_failure():
    dataset,values=outcome_fixture();values[0]['converted']=None
    result=a.uploaded_experiment(dataset,values,a.source_meta(dataset,Filters()))
    assert result['arms'][0]['conversion_rate'] is None and result['arms'][0]['applications'] is None


def test_workspace_owner_is_bound_in_lookup(monkeypatch):
    calls=[];identifier=str(uuid.uuid4())
    monkeypatch.setattr(u,'workspace_hash',lambda token:'ownerA' if token=='tokenA' else 'ownerB')
    def lookup(sql,params):
        calls.append((sql,params))
        return {'dataset_id':identifier,'name':'private','metadata':{}} if params[1]=='ownerA' else {}
    monkeypatch.setattr(u,'row',lookup)
    assert u.get_dataset(identifier,'tokenA')['name']=='private'
    with pytest.raises(u.UploadError,match='this workspace'):u.get_dataset(identifier,'tokenB')
    assert all('workspace_hash=%s' in sql for sql,_ in calls)
    assert calls[1][1]==(identifier,'ownerB')


@pytest.mark.parametrize('quota',[
 {'own':0,'total':0,'own_bytes':u.WORKSPACE_PAYLOAD_BUDGET,'total_bytes':0},
 {'own':0,'total':0,'own_bytes':0,'total_bytes':u.GLOBAL_PAYLOAD_BUDGET},
 {'own':10,'total':10,'own_bytes':0,'total_bytes':0},
])
def test_storage_quota_locked_before_insert(monkeypatch,quota):
    calls=[]
    class Result:
        def fetchone(self):return quota
    class Conn:
        def execute(self,sql,*args):calls.append(sql);return Result()
    @contextmanager
    def connect(**kwargs):yield Conn()
    monkeypatch.setattr(u,'connection',connect);monkeypatch.setattr(u,'workspace_hash',lambda _:'owner')
    dataset,records=fixture([{'job_title':'Analyst'}]);raw=[r['row_data'] for r in records]
    prepared=u.prepare_rows(dataset['columns'],raw)
    with pytest.raises(u.UploadError,match='quota'):u.store_dataset('token','x.csv',b'x',dataset['columns'],raw,prepared,{})
    assert 'pg_advisory_xact_lock' in calls[0]
    assert not any('INSERT' in c for c in calls)


def test_model_mapping_suggestions_are_allowlisted(monkeypatch):
    from api.analytics import providers
    monkeypatch.setenv('OLLAMA_API_KEY','test-key')
    async def complete(*args,**kwargs):
        assert kwargs['max_tokens']>=800
        assert 'supported_canonical_fields' in args[2]
        return {'content':json.dumps({'description':'Catalog','grain':'job','suggested_dimensions':['market','execute_sql'],'mapping_suggestions':{'job_title':'Title','sql':'DROP','market':'missing'}}),'latency_ms':50,'cost_usd':None}
    monkeypatch.setattr(providers,'complete',complete)
    result=asyncio.run(u.semantic_metadata(['Title'],[{'Title':'Analyst'}],{}))
    assert result['mode']=='provider' and result['mapping_suggestions']=={'job_title':'Title'}
    assert result['suggested_dimensions']==['market'] and result['latency_ms']==50


def test_csv_extreme_cell_is_validation_error():
    with pytest.raises(u.UploadError,match='oversized'):u.parse_file(b'job_title\n'+b'x'*200000,'bad.csv')


def test_invalid_xml_is_validation_error():
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w') as z:z.writestr('xl/workbook.xml',b'<!DOCTYPE x [<!ENTITY a "x">]><x>&a;</x>')
    with pytest.raises(u.UploadError,match='declarations'):u.parse_file(buf.getvalue(),'entity.xlsx')


def test_user_outcome_overview_and_segments_deduplicate(monkeypatch):
    dataset,values=outcome_fixture(False)
    records=[{'row_number':i+1,'row_data':v,'normalized_data':v} for i,v in enumerate(values+values)]
    monkeypatch.setattr(a,'dataset_rows',lambda *_:(dataset,records))
    result=a.analyze_dataset('overview',dataset['dataset_id'],'token',Filters())
    assert next(k['value'] for k in result['kpis'] if k['key']=='exposed_users')==300
    assert next(k['value'] for k in result['kpis'] if k['key']=='applications')==45
    assert result['product_metrics']['dau'] is None
    segments=a.analyze_dataset('segments',dataset['dataset_id'],'token',Filters(),dimension='variant')
    assert segments['metric']=='conversion_rate' and all(r['sessions'] is None for r in segments['rows'])
    assert {r['name']:r['conversion_rate'] for r in segments['rows']}=={'control':10,'gpt4o':20,'ollama':15}


def test_gemma_fenced_json_metadata(monkeypatch):
    from api.analytics import providers
    monkeypatch.setenv('OLLAMA_API_KEY','test-key')
    async def complete(*args,**kwargs):return {'content':'```json\n{"description":"Catalog","mapping_suggestions":{"job_title":"Title"}}\n```','latency_ms':100,'cost_usd':None}
    monkeypatch.setattr(providers,'complete',complete)
    result=asyncio.run(u.semantic_metadata(['Title'],[{'Title':'Analyst'}],{}))
    assert result['mode']=='provider' and result['mapping_suggestions']=={'job_title':'Title'}


def test_uploads_do_not_invent_allocation_plan_or_observation_dates():
    dataset,values=outcome_fixture(True)
    result=a.uploaded_experiment(dataset,values,a.source_meta(dataset,Filters()))
    assert result['srm']['status']=='planned_allocation_unknown' and result['srm']['p_value'] is None
    dataset['mapping'].pop('date')
    result=a.uploaded_experiment(dataset,values,a.source_meta(dataset,Filters()))
    assert result['status']=='descriptive_only'
    assert result['meta']['date_filter_applied'] is False


def test_event_cohort_filter_preserves_outcomes_with_sparse_properties():
    records=[{'row_number':1,'normalized_data':{'date':'2026-10-04','session_id':'s','user_id':'u','market':'Germany','event_name':'search_started'}},{'row_number':2,'normalized_data':{'date':'2026-10-04','session_id':'s','event_name':'application_submitted'}}]
    selected=a.select_rows(records,Filters(market='Germany'),kind='events')
    assert len(selected)==2
    sessions=a.behavioral_sessions(selected,'events')
    assert sessions[0]['application_submitted'] is True and sessions[0]['market']=='Germany'


def test_uploaded_ai_user_outcomes_uses_user_numeric_evidence(monkeypatch):
    dataset,values=outcome_fixture(False)
    records=[{'row_number':i+1,'row_data':v,'normalized_data':v} for i,v in enumerate(values)]
    monkeypatch.setattr(a,'dataset_rows',lambda *_:(dataset,records))
    result=a.uploaded_evidence(dataset['dataset_id'],'token','Visualize conversion by variant',Filters())
    assert 'distinct exposed uploaded users' in result['answer']
    assert {r['name']:r['conversion_rate'] for r in result['chart']['data']}=={'control':10,'gpt4o':20,'ollama':15}
